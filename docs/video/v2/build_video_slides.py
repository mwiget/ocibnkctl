#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx"]
# ///
"""
Render the slides for the v2 (BNK 2.4.0) demo video on the F5 FY26 template.

Style mirrors tmm-lb-nico/docs/build_nico_teaser_deck.py — same palette, same
title/chip/card helpers — so the video opens and closes the way that deck does.

The template is NOT vendored. It is an internal F5 asset and this repo is
public: it is read from --template, and the built .pptx plus the rendered PNGs
land in build/, which is git-ignored.

    uv run docs/video/v2/build_video_slides.py \
        --template ~/git/F5-Template-Starter-FY26-Q2.potx   # -> build/slides/*.png
"""

import argparse
import os
import shutil
import subprocess
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))

BG_DARK   = RGBColor(0x0A, 0x0A, 0x1A)
BG_SLIDE  = RGBColor(0x0F, 0x14, 0x2A)
F5_BLUE   = RGBColor(0x2E, 0x6E, 0xE6)
F5_PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
F5_RED    = RGBColor(0xE4, 0x00, 0x2B)
F5_ORANGE = RGBColor(0xF5, 0x9E, 0x0B)
F5_GREEN  = RGBColor(0x10, 0xB9, 0x81)
F5_CYAN   = RGBColor(0x06, 0xB6, 0xD4)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GY  = RGBColor(0xBB, 0xBB, 0xCC)
MID_GY    = RGBColor(0x88, 0x88, 0x99)
CARD_BG   = RGBColor(0x16, 0x1E, 0x38)
CARD_DARK = RGBColor(0x0D, 0x11, 0x22)
MONO      = "Consolas"
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

REPO = "github.com/mwiget/ocibnkctl"


def load_template(path):
    """.potx -> .pptx content-type shim, then strip the example slides."""
    if path.lower().endswith(".potx"):
        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False).name
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                zout.writestr(item, data)
        prs = Presentation(tmp)
        os.unlink(tmp)
    else:
        prs = Presentation(path)
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        try:
            prs.part.drop_rel(sid.get(rel_ns))
        except KeyError:
            pass
        lst.remove(sid)
    return prs


# ── helpers (from the teaser deck) ────────────────────────────────────────────
def _bg(s, c):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = c

def _rect(s, l, t, w, h, fc, bc=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fc
    if bc: sh.line.color.rgb = bc; sh.line.width = Pt(1.5)
    else:  sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def _rrect(s, l, t, w, h, fc, bc=None):
    sh = _rect(s, l, t, w, h, fc, bc, MSO_SHAPE.ROUNDED_RECTANGLE)
    sh.adjustments[0] = min(0.5, Inches(0.16) / min(w, h))  # same corner radius at any size
    return sh

def _tb(s, l, t, w, h):
    return s.shapes.add_textbox(l, t, w, h)

def _txt(tf, text, sz=18, c=WHITE, b=False, al=PP_ALIGN.LEFT, fn="Calibri"):
    tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(sz); p.font.color.rgb = c
    p.font.bold = b; p.font.name = fn; p.alignment = al
    return p

def _p(tf, text, sz=18, c=WHITE, b=False, fn="Calibri", sb=Pt(4), sa=Pt(2)):
    p = tf.add_paragraph()
    p.text = text; p.font.size = Pt(sz); p.font.color.rgb = c
    p.font.bold = b; p.font.name = fn; p.space_before = sb; p.space_after = sa
    return p

def _run(p, text, sz, c, b=False, fn="Calibri"):
    r = p.add_run(); r.text = text
    r.font.size = Pt(sz); r.font.color.rgb = c; r.font.bold = b; r.font.name = fn
    return r

def _chip_label(s, label, col):
    _rrect(s, Inches(10.55), Inches(0.3), Inches(2.2), Inches(0.36), CARD_DARK, col)
    tb = _tb(s, Inches(10.55), Inches(0.345), Inches(2.2), Inches(0.3))
    _txt(tb.text_frame, label, sz=11, c=col, b=True, al=PP_ALIGN.CENTER)

def _slide_title(s, title, subtitle, chip=None):
    tb = _tb(s, Inches(0.7), Inches(0.22), Inches(9.9), Inches(0.6))
    _txt(tb.text_frame, title, sz=30, b=True)
    tb = _tb(s, Inches(0.72), Inches(0.82), Inches(12), Inches(0.4))
    _txt(tb.text_frame, subtitle, sz=15, c=F5_CYAN)
    if chip:
        _chip_label(s, *chip)

def _box(s, x, y, w, h, fill, border, label, sub=None, lsz=14, fn=MONO):
    sh = _rrect(s, Inches(x), Inches(y), Inches(w), Inches(h), fill, border)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Pt(4))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, label, lsz, WHITE, b=True, fn=fn)
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        _run(p2, sub, 10, LIGHT_GY)
    return sh

def _card(s, x, y, w, h, title, col, lines, tsz=15, bsz=12.5):
    _rrect(s, Inches(x), Inches(y), Inches(w), Inches(h), CARD_BG, col)
    tb = _tb(s, Inches(x + 0.2), Inches(y + 0.12), Inches(w - 0.4), Inches(0.4))
    _txt(tb.text_frame, title, sz=tsz, c=col, b=True)
    tb = _tb(s, Inches(x + 0.2), Inches(y + 0.55), Inches(w - 0.4), Inches(h - 0.65))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            _txt(tf, line, sz=bsz, c=LIGHT_GY)
        else:
            _p(tf, line, sz=bsz, c=LIGHT_GY, sb=Pt(3), sa=Pt(0))


def build(prs):
    layouts = {l.name: l for l in prs.slide_layouts}
    title_layout = layouts.get("Title Slide: Black Solid")
    blank_layout = layouts.get("Blank Layout Black")

    def title_slide():
        s = prs.slides.add_slide(title_layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        _bg(s, BG_DARK)
        _rect(s, 0, 0, SLIDE_W, Inches(0.06), F5_RED)
        return s

    def content_slide():
        s = prs.slides.add_slide(blank_layout)
        _bg(s, BG_SLIDE)
        _rect(s, 0, 0, SLIDE_W, Inches(0.04), F5_BLUE)
        return s

    out = []

    # ── title ────────────────────────────────────────────────────────────────
    s = title_slide()
    tb = _tb(s, Inches(1.2), Inches(1.1), Inches(11.7), Inches(1.9))
    _txt(tb.text_frame, "F5 BIG-IP Next for Kubernetes 2.4.0", sz=40, b=True)
    _p(tb.text_frame, "on plain Docker", sz=40, b=True, sb=Pt(0), sa=Pt(0))
    tb = _tb(s, Inches(1.2), Inches(3.15), Inches(11), Inches(0.8))
    _txt(tb.text_frame, "Deployed, tested and torn down by an AI agent on a local model",
         sz=20, c=F5_CYAN, b=True)
    _rect(s, Inches(1.2), Inches(4.05), Inches(3), Inches(0.04), F5_RED)
    tb = _tb(s, Inches(1.2), Inches(4.35), Inches(11), Inches(1.2))
    _txt(tb.text_frame, "ocibnkctl · one binary · k3s nodes as containers · 15 F5 how-tos verified",
         sz=17, c=MID_GY, fn=MONO)
    tb = _tb(s, Inches(1.2), Inches(6.35), Inches(11), Inches(0.5))
    _txt(tb.text_frame, REPO, sz=15, c=LIGHT_GY, fn=MONO)
    out.append("title")

    # ── why: BNK on one host ─────────────────────────────────────────────────
    s = content_slide()
    _slide_title(s, "BNK on one host",
                 "A full BNK 2.4.0 stack — control plane, TMM, BGP edge — as containers on Docker or podman",
                 ("ONE HOST", F5_BLUE))
    # host frame
    _rrect(s, Inches(0.7), Inches(1.5), Inches(7.3), Inches(5.3), CARD_DARK, MID_GY)
    tb = _tb(s, Inches(0.9), Inches(1.58), Inches(7), Inches(0.35))
    _txt(tb.text_frame, "host  ·  Docker / podman  ·  nothing else installed", sz=12, c=MID_GY, fn=MONO)
    _box(s, 1.0, 2.05, 3.2, 1.5, CARD_BG, F5_BLUE, "k3s server",
         "control node · tainted NoSchedule\nAPI · Calico · cert-manager")
    _box(s, 4.5, 2.05, 3.2, 1.5, CARD_BG, F5_RED, "k3s agent",
         "labelled app=f5-tmm\nFLO · CNEInstance · TMM (demo mode)")
    _box(s, 1.0, 3.85, 3.2, 1.2, CARD_BG, F5_GREEN, "edge FRR",
         "BGP peer to TMM — VIPs as /32s")
    _box(s, 4.5, 3.85, 3.2, 1.2, CARD_BG, F5_GREEN, "origins",
         "HTTP · gRPC · TCP · UDP · AI backends")
    _box(s, 1.0, 5.35, 6.7, 1.2, CODE_BG := RGBColor(0x08, 0x0C, 0x18), F5_CYAN,
         "validate → cluster up → prereqs → flo → [shrink] → cne",
         "ocibnkctl e2e — every phase idempotent and resume-safe; poc.yaml is the source of truth",
         lsz=13)
    # claims
    _card(s, 8.3, 1.5, 4.35, 1.6, "No kind, no k3d, no DPU", F5_BLUE,
          ["k3s runs straight on the container runtime.",
           "TMM in demo mode: virtio, no SR-IOV."])
    _card(s, 8.3, 3.25, 4.35, 1.6, "Laptop-sized", F5_ORANGE,
          ["10 cores is the floor.",
           "Tight on memory? deploy shrink engages by itself."])
    _card(s, 8.3, 5.0, 4.35, 1.8, "Down to a Raspberry Pi 5", F5_GREEN,
          ["4 cores, arm64: host_profile: small",
           "plus the automatic shrink phase.",
           "(measured on BNK 2.3.0)"])
    out.append("one-host")

    # ── what the scenarios prove ─────────────────────────────────────────────
    s = content_slide()
    _slide_title(s, "What the scenarios prove",
                 "15 scenarios, each one an F5 how-to article — applied, asserted, reported",
                 ("PROOF", F5_GREEN))
    _card(s, 0.7, 1.5, 6.0, 2.75, "Gateway API traffic", F5_BLUE,
          ["HTTP routing (how-to #8) · external pool (#10)",
           "TCP and UDP L4 load balancing",
           "gRPC load balancing",
           "Proxy Protocol iRule on L4 (#9)",
           "Dynamic IP allocation (FIC)"])
    _card(s, 6.95, 1.5, 5.7, 2.75, "BGP at the edge", F5_CYAN,
          ["TMM ZebOS ⇄ external FRR peering",
           "BGP anycast, all-active data plane (#3)",
           "VIPs reach the client as real routes —",
           "no port-forward, no shortcut"])
    _card(s, 0.7, 4.45, 6.0, 2.35, "AI gateway", F5_PURPLE,
          ["Token counting and enforcement (#6)",
           "Token counting via iRule + DSSM",
           "Semantic model caching (#7)"])
    _card(s, 6.95, 4.45, 5.7, 2.35, "Operations and security", F5_ORANGE,
          ["Restricted CWC admin access (#1)",
           "Cluster-wide watch (#2)",
           "Core file collection (#4)"])
    out.append("scenarios")

    # ── close ────────────────────────────────────────────────────────────────
    s = title_slide()
    tb = _tb(s, Inches(1.2), Inches(1.4), Inches(11.7), Inches(1.2))
    _txt(tb.text_frame, REPO, sz=40, b=True, fn=MONO)
    tb = _tb(s, Inches(1.2), Inches(2.7), Inches(11), Inches(0.8))
    _txt(tb.text_frame, "Single-binary download for Linux and macOS — releases/latest",
         sz=21, c=F5_CYAN, b=True)
    _rect(s, Inches(1.2), Inches(3.65), Inches(3), Inches(0.04), F5_RED)
    tb = _tb(s, Inches(1.2), Inches(3.95), Inches(11), Inches(1.6))
    tf = tb.text_frame
    _txt(tf, "bnkscope     github.com/mwiget/bnkscope     — see what the cluster is doing",
         sz=16, c=MID_GY, fn=MONO)
    _p(tf, "regcachectl  github.com/mwiget/regcachectl  — warm image pulls between runs",
       sz=16, c=MID_GY, fn=MONO)
    out.append("close")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", default=os.path.expanduser("~/git/F5-Template-Starter-FY26-Q2.potx"))
    ap.add_argument("--out", default=os.path.join(HERE, "build", "slides"))
    ap.add_argument("--dpi", type=int, default=144)  # 13.333 in × 144 = 1920 px
    args = ap.parse_args()

    prs = load_template(args.template)
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    names = build(prs)

    os.makedirs(args.out, exist_ok=True)
    pptx = os.path.join(os.path.dirname(args.out), "ocibnkctl-video-slides.pptx")
    prs.save(pptx)
    with tempfile.TemporaryDirectory() as work:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", work, pptx],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = os.path.join(work, os.path.basename(pptx)[:-5] + ".pdf")
        subprocess.run(["pdftoppm", "-r", str(args.dpi), "-png", pdf, os.path.join(work, "s")], check=True)
        pages = sorted(f for f in os.listdir(work) if f.startswith("s") and f.endswith(".png"))
        for name, page in zip(names, pages):
            shutil.move(os.path.join(work, page), os.path.join(args.out, f"{name}.png"))
    print("slides:", ", ".join(names), "->", args.out)


if __name__ == "__main__":
    main()
